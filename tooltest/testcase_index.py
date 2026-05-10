from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Set

import openpyxl


ROOT_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = ROOT_DIR.parent
WORKBOOK_PATH = WORKSPACE_DIR / "13_system_test.xlsx"
SELENIUM_DIR = ROOT_DIR / "selenium"
POSTMAN_COLLECTION = ROOT_DIR / "postman" / "online_learning_collection.json"
JMETER_PLAN = ROOT_DIR / "jmeter" / "online_learning_perf_test.jmx"
TEST_DATA_DIR = ROOT_DIR / "test_data"
OUTPUT_PATH = ROOT_DIR / "TRACEABILITY.md"

TC_PATTERN = re.compile(r"TC-[A-Z]+-\d{3}[A-Z]*")


def load_workbook_cases() -> List[Dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    cases = []
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(values_only=True):
            if len(row) <= 1:
                continue
            tc_id = str(row[1]).strip() if row[1] else ""
            if not tc_id.startswith("TC-"):
                continue
            cases.append(
                {
                    "tc_id": tc_id,
                    "sheet": sheet_name,
                    "scenario": str(row[4]).strip() if row[4] else "",
                    "manual_status": str(row[9]).strip() if row[9] else "",
                }
            )
    return cases


def scan_python_files(directory: Path) -> Dict[str, Set[str]]:
    found: Dict[str, Set[str]] = defaultdict(set)
    for path in directory.glob("*.py"):
        content = path.read_text(encoding="utf-8", errors="ignore")
        for tc_id in sorted(set(TC_PATTERN.findall(content))):
            found[tc_id].add(path.name)
    return found


def scan_jmeter_plan(path: Path) -> Dict[str, Set[str]]:
    content = path.read_text(encoding="utf-8", errors="ignore")
    found: Dict[str, Set[str]] = defaultdict(set)
    for tc_id in sorted(set(TC_PATTERN.findall(content))):
        found[tc_id].add(path.name)
    return found


def walk_postman_items(items: Iterable[Dict], found: Dict[str, Set[str]]):
    for item in items:
        name = item.get("name", "")
        matches = set(TC_PATTERN.findall(name))
        for tc_id in matches:
            found[tc_id].add("online_learning_collection.json")
        if "item" in item:
            walk_postman_items(item["item"], found)


def scan_postman_collection(path: Path) -> Dict[str, Set[str]]:
    if not path.exists():
        return {}
    collection = json.loads(path.read_text(encoding="utf-8"))
    found: Dict[str, Set[str]] = defaultdict(set)
    walk_postman_items(collection.get("item", []), found)
    return found


def scan_test_data(directory: Path) -> Dict[str, Set[str]]:
    found: Dict[str, Set[str]] = defaultdict(set)
    for path in directory.glob("*.csv"):
        content = path.read_text(encoding="utf-8-sig", errors="ignore")
        for tc_id in sorted(set(TC_PATTERN.findall(content))):
            found[tc_id].add(path.name)
    return found


def format_files(files: Set[str]) -> str:
    return ", ".join(sorted(files)) if files else "-"


def build_markdown() -> str:
    cases = load_workbook_cases()
    selenium_hits = scan_python_files(SELENIUM_DIR)
    postman_hits = scan_postman_collection(POSTMAN_COLLECTION)
    jmeter_hits = scan_jmeter_plan(JMETER_PLAN)
    data_hits = scan_test_data(TEST_DATA_DIR)

    sheet_counter = Counter(case["sheet"] for case in cases)
    selenium_count = sum(1 for case in cases if case["tc_id"] in selenium_hits)
    postman_count = sum(1 for case in cases if case["tc_id"] in postman_hits)
    jmeter_count = sum(1 for case in cases if case["tc_id"] in jmeter_hits)

    lines = [
        "# Tooltest Traceability",
        "",
        f"- Source workbook: `{WORKBOOK_PATH.name}`",
        f"- Total manual system test cases: **{len(cases)}**",
        f"- Selenium references found: **{selenium_count}**",
        f"- Postman references found: **{postman_count}**",
        f"- JMeter references found: **{jmeter_count}**",
        "",
        "## Workbook Sheets",
        "",
        "| Sheet | Cases |",
        "|---|---:|",
    ]

    for sheet_name, count in sorted(sheet_counter.items()):
        lines.append(f"| {sheet_name} | {count} |")

    lines.extend(
        [
            "",
            "## Case Index",
            "",
            "| TC ID | Sheet | Manual Status | Selenium | Postman | JMeter | Test Data | Scenario |",
            "|---|---|---|---|---|---|---|---|",
        ]
    )

    for case in sorted(cases, key=lambda item: item["tc_id"]):
        tc_id = case["tc_id"]
        scenario = case["scenario"].replace("\n", " ").strip()
        lines.append(
            "| "
            + " | ".join(
                [
                    tc_id,
                    case["sheet"],
                    case["manual_status"] or "-",
                    format_files(selenium_hits.get(tc_id, set())),
                    format_files(postman_hits.get(tc_id, set())),
                    format_files(jmeter_hits.get(tc_id, set())),
                    format_files(data_hits.get(tc_id, set())),
                    scenario or "-",
                ]
            )
            + " |"
        )

    missing_selenium = [case["tc_id"] for case in cases if case["tc_id"].startswith(("TC-AUTH-", "TC-COURSE-", "TC-ORDER-", "TC-STUDY-")) and case["tc_id"] not in selenium_hits]
    lines.extend(
        [
            "",
            "## Missing Selenium Coverage In UI Scope",
            "",
            ", ".join(missing_selenium) if missing_selenium else "None",
            "",
        ]
    )
    return "\n".join(lines)


def main():
    markdown = build_markdown()
    OUTPUT_PATH.write_text(markdown, encoding="utf-8")
    print(f"Saved traceability report to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
