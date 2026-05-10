from __future__ import annotations

import csv
from functools import lru_cache
from pathlib import Path
from typing import Dict, Iterable, List

import openpyxl


WORKSPACE_DIR = Path(__file__).resolve().parents[2]
TOOLTEST_DIR = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = WORKSPACE_DIR / "13_system_test.xlsx"
TEST_DATA_DIR = TOOLTEST_DIR / "test_data"

WORKBOOK_COLUMNS = {
    "tc_id": 1,
    "feature": 2,
    "precondition": 3,
    "scenario": 4,
    "steps": 5,
    "input_data": 6,
    "expected_result": 7,
    "manual_actual": 8,
    "manual_status": 9,
}


def _clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "none" else text


def _row_to_case(sheet_name: str, row: tuple) -> Dict[str, str]:
    case = {"sheet": sheet_name, "source": WORKBOOK_PATH.name}
    for field, index in WORKBOOK_COLUMNS.items():
        case[field] = _clean(row[index]) if len(row) > index else ""
    return case


def _csv_row_to_text(row: Dict[str, str]) -> str:
    parts = []
    for key, value in row.items():
        if key == "tc_id":
            continue
        cleaned = _clean(value)
        if cleaned:
            parts.append(f"{key}={cleaned}")
    return "; ".join(parts)


@lru_cache(maxsize=1)
def load_workbook_cases() -> Dict[str, Dict[str, str]]:
    if not WORKBOOK_PATH.exists():
        return {}

    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    cases: Dict[str, Dict[str, str]] = {}
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(values_only=True):
            if len(row) <= 1:
                continue
            tc_id = _clean(row[WORKBOOK_COLUMNS["tc_id"]])
            if not tc_id.startswith("TC-"):
                continue
            cases[tc_id] = _row_to_case(sheet_name, row)
    return cases


@lru_cache(maxsize=1)
def load_csv_case_data() -> Dict[str, Dict[str, str]]:
    data: Dict[str, Dict[str, str]] = {}
    if not TEST_DATA_DIR.exists():
        return data

    for csv_path in sorted(TEST_DATA_DIR.glob("*.csv")):
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                tc_id = _clean(row.get("tc_id"))
                if not tc_id:
                    continue
                cleaned_row = {key: _clean(value) for key, value in row.items()}
                cleaned_row["csv_file"] = csv_path.name
                cleaned_row["csv_summary"] = _csv_row_to_text(cleaned_row)
                data[tc_id] = cleaned_row
    return data


def get_test_case(tc_id: str) -> Dict[str, str]:
    workbook_case = dict(load_workbook_cases().get(tc_id, {}))
    csv_case = load_csv_case_data().get(tc_id, {})
    if csv_case:
        workbook_case["csv_file"] = csv_case.get("csv_file", "")
        workbook_case["csv_summary"] = csv_case.get("csv_summary", "")
    return workbook_case


def get_scope_cases(prefixes: Iterable[str]) -> List[Dict[str, str]]:
    matched = []
    for tc_id, case in load_workbook_cases().items():
        if any(tc_id.startswith(prefix) for prefix in prefixes):
            matched.append(case)
    return sorted(matched, key=lambda item: item["tc_id"])

